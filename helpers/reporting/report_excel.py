# -*- coding: utf-8 -*-
"""
Excel Report Generator.
Builds a professional, multi-sheet Excel workbook using openpyxl.
Includes Summary, Reference Profile, Layer Overview, Layer Diffs, Geometry Comparison,
Property Comparison, Unexpected Features, Closure Issues, Count Diffs, Geometry QC, and Issue Details.
"""

import os
from typing import Optional, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models.qc_result import QCResult, LayerStatus
from ..models.qc_issue import Severity, CheckID


def create_excel_report(qc_result: QCResult, output_path: str) -> str:
    """
    Generates a complete multi-sheet Excel report from a QCResult.
    Returns the absolute path to the generated Excel file.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling palettes
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)

    fill_pass = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # soft green
    fill_error = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # soft red
    fill_warning = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # soft yellow
    fill_info = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")  # soft blue

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    def style_header_row(ws, row_idx, num_cols):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def autofit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len and "\n" not in val:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ==========================================
    # 1. Sheet: Summary
    # ==========================================
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["A1"] = "CAD STANDARDS & QC ANALYZER REPORT"
    ws_sum["A1"].font = title_font
    ws_sum["A2"] = f"Execution Date: {qc_result.execution_date}"
    ws_sum["A2"].font = subtitle_font

    ws_sum["A4"] = "EXECUTIVE METRIC"
    ws_sum["B4"] = "VALUE"
    style_header_row(ws_sum, 4, 2)

    sum_data = [
        ("Reference Source", qc_result.reference_source),
        ("Received Source", qc_result.received_source),
        ("Spatial Reference (Reference)", qc_result.spatial_reference_ref),
        ("Spatial Reference (Received)", qc_result.spatial_reference_rec),
        ("Total Reference Layers", qc_result.total_reference_layers()),
        ("Total Received Layers", qc_result.total_received_layers()),
        ("Matching Layers", len(qc_result.matching_layers)),
        ("Missing Layers", len(qc_result.missing_layers)),
        ("Extra Layers", len(qc_result.extra_layers)),
        ("Passed Layers", qc_result.passed_layers_count()),
        ("Layers with Errors", qc_result.error_layers_count()),
        ("Layers with Warnings", qc_result.warning_layers_count()),
        ("Reference Comparison Issues", len(qc_result.reference_comparison_issues())),
        ("Geometry QC Issues", len(qc_result.geometry_qc_issues())),
        ("Total QC Issues", len(qc_result.issues)),
        ("Total Errors", len(qc_result.issues_by_severity(Severity.ERROR))),
        ("Total Warnings", len(qc_result.issues_by_severity(Severity.WARNING))),
    ]

    r = 5
    for label, val in sum_data:
        c1 = ws_sum.cell(row=r, column=1, value=label)
        c2 = ws_sum.cell(row=r, column=2, value=val)
        c1.font = bold_font
        c2.font = normal_font
        c1.border = thin_border
        c2.border = thin_border
        r += 1

    autofit_columns(ws_sum)

    # ==========================================
    # 2. Sheet: Reference Profile
    # ==========================================
    ws_ref = wb.create_sheet(title="Reference Profile")
    ws_ref.views.sheetView[0].showGridLines = True

    ws_ref["A1"] = "REFERENCE PROFILE (Observed Standards from Original CAD)"
    ws_ref["A1"].font = title_font

    ref_headers = [
        "Layer Name",
        "Feature Count",
        "Dominant Geometry",
        "Dominant %",
        "Composition",
        "Closure Required",
        "Closure %",
        "Dominant Color",
        "Dominant Linetype",
        "Dominant Lineweight",
    ]
    for c_idx, h in enumerate(ref_headers, 1):
        ws_ref.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_ref, 3, len(ref_headers))

    r = 4
    if qc_result.reference_profile:
        for l_name, lp in sorted(qc_result.reference_profile.layers.items()):
            row_vals = [
                lp.layer_name,
                lp.feature_count,
                lp.dominant_geometry or "None",
                f"{lp.dominant_percentage * 100:.1f}%",
                "MIXED" if lp.is_mixed_geometry else "CONSISTENT",
                "Yes" if lp.closure_required else "No",
                f"{lp.closure_percentage * 100:.1f}%" if lp.closure_required else "N/A",
                lp.dominant_color() or "N/A",
                lp.dominant_linetype() or "N/A",
                lp.dominant_lineweight() or "N/A",
            ]
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws_ref.cell(row=r, column=c_idx, value=val)
                cell.font = normal_font
                cell.border = thin_border
            r += 1

    autofit_columns(ws_ref)

    # ==========================================
    # 3. Sheet: Layer Overview
    # ==========================================
    ws_ov = wb.create_sheet(title="Layer Overview")
    ws_ov.views.sheetView[0].showGridLines = True

    ws_ov["A1"] = "LAYER OVERVIEW & COMPLIANCE STATUS"
    ws_ov["A1"].font = title_font

    ov_headers = [
        "Layer Name",
        "Received Count",
        "Expected Geometry",
        "Received Geometry",
        "Geometry Status",
        "Color",
        "Lineweight",
        "Linetype",
        "Closure",
        "Status",
        "Reasons",
    ]
    for c_idx, h in enumerate(ov_headers, 1):
        ws_ov.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_ov, 3, len(ov_headers))

    r = 4
    for l_name, l_stat in sorted(qc_result.layer_statuses.items()):
        row_vals = [
            l_stat.layer_name,
            l_stat.rec_count,
            l_stat.expected_geom,
            l_stat.received_geom,
            l_stat.geom_status,
            l_stat.color_status,
            l_stat.lineweight_status,
            l_stat.linetype_status,
            l_stat.closure_status,
            l_stat.status,
            "; ".join(l_stat.reasons),
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_ov.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border

            # Highlight status column
            if c_idx == 10:
                if val == LayerStatus.PASS:
                    cell.fill = fill_pass
                elif val in (LayerStatus.ERROR, LayerStatus.MISSING):
                    cell.fill = fill_error
                elif val in (LayerStatus.WARNING, LayerStatus.REVIEW):
                    cell.fill = fill_warning
                elif val == LayerStatus.EXTRA:
                    cell.fill = fill_info
        r += 1

    autofit_columns(ws_ov)

    # ==========================================
    # 4. Sheet: Layer Differences
    # ==========================================
    ws_diff = wb.create_sheet(title="Layer Differences")
    ws_diff.views.sheetView[0].showGridLines = True

    ws_diff["A1"] = "LAYER DIFFERENCES (Missing & Extra Layers)"
    ws_diff["A1"].font = title_font

    diff_headers = ["Layer Name", "Classification", "Received Count", "Severity", "Details"]
    for c_idx, h in enumerate(diff_headers, 1):
        ws_diff.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_diff, 3, len(diff_headers))

    r = 4
    for m in qc_result.missing_layers:
        lp = qc_result.reference_profile.layers.get(m) if qc_result.reference_profile else None
        l_name = lp.layer_name if lp else m
        vals = [l_name, "MISSING LAYER", 0, "ERROR", "Layer exists in Reference but missing from Received"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_diff.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 2:
                cell.fill = fill_error
        r += 1

    for e in qc_result.extra_layers:
        lp = qc_result.received_profile.layers.get(e) if qc_result.received_profile else None
        l_name = lp.layer_name if lp else e
        f_cnt = lp.feature_count if lp else 0
        vals = [l_name, "EXTRA LAYER", f_cnt, "INFO", "Layer exists in Received but not in Reference"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_diff.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 2:
                cell.fill = fill_info
        r += 1

    autofit_columns(ws_diff)

    # ==========================================
    # 5. Sheet: Geometry Comparison
    # ==========================================
    ws_geom = wb.create_sheet(title="Geometry Comparison")
    ws_geom.views.sheetView[0].showGridLines = True

    ws_geom["A1"] = "GEOMETRY DISTRIBUTION COMPARISON"
    ws_geom["A1"].font = title_font

    geom_headers = ["Layer Name", "Geometry Category", "Received Count", "Status"]
    for c_idx, h in enumerate(geom_headers, 1):
        ws_geom.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_geom, 3, len(geom_headers))

    r = 4
    if qc_result.reference_profile and qc_result.received_profile:
        for key in qc_result.matching_layers:
            lp_ref = qc_result.reference_profile.layers.get(key)
            lp_rec = qc_result.received_profile.layers.get(key)
            if not lp_ref or not lp_rec:
                continue
            all_cats = sorted(list(set(lp_ref.geometry_distribution.keys()).union(set(lp_rec.geometry_distribution.keys()))))
            for cat in all_cats:
                c_ref = lp_ref.geometry_distribution.get(cat, 0)
                c_rec = lp_rec.geometry_distribution.get(cat, 0)
                status = "MATCH"
                if c_ref == 0 and c_rec > 0:
                    status = "UNEXPECTED"
                elif c_ref > 0 and c_rec == 0:
                    status = "MISSING TYPE"

                vals = [lp_ref.layer_name, cat, c_rec, status]
                for c_idx, val in enumerate(vals, 1):
                    cell = ws_geom.cell(row=r, column=c_idx, value=val)
                    cell.font = normal_font
                    cell.border = thin_border
                    if c_idx == 4:
                        if status == "UNEXPECTED":
                            cell.fill = fill_error
                        elif status == "MISSING TYPE":
                            cell.fill = fill_warning
                        elif status == "MATCH":
                            cell.fill = fill_pass
                r += 1

    autofit_columns(ws_geom)

    # ==========================================
    # 6. Sheet: Property Comparison
    # ==========================================
    ws_prop = wb.create_sheet(title="Property Comparison")
    ws_prop.views.sheetView[0].showGridLines = True

    ws_prop["A1"] = "CAD PROPERTY COMPARISON (Color, Linetype, Lineweight)"
    ws_prop["A1"].font = title_font

    prop_headers = ["Layer Name", "Property", "Expected (Reference)", "Actual (Received)", "Status", "Details"]
    for c_idx, h in enumerate(prop_headers, 1):
        ws_prop.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_prop, 3, len(prop_headers))

    r = 4
    prop_issues = [i for i in qc_result.issues if i.check_id in (CheckID.CHK_COLOR, CheckID.CHK_LINEWEIGHT, CheckID.CHK_LINETYPE)]
    for pi in prop_issues:
        vals = [pi.layer_name, pi.property_name, pi.expected_value, pi.actual_value, pi.severity, pi.details]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_prop.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 5:
                cell.fill = fill_warning if pi.severity == Severity.WARNING else fill_error
        r += 1

    autofit_columns(ws_prop)

    # ==========================================
    # 7. Sheet: Unexpected Features
    # ==========================================
    ws_unexp = wb.create_sheet(title="Unexpected Features")
    ws_unexp.views.sheetView[0].showGridLines = True

    ws_unexp["A1"] = "UNEXPECTED FEATURES INSIDE LAYERS"
    ws_unexp["A1"].font = title_font

    unexp_headers = ["Layer Name", "Expected Geometry", "Unexpected Geometry", "Affected Features", "Details"]
    for c_idx, h in enumerate(unexp_headers, 1):
        ws_unexp.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_unexp, 3, len(unexp_headers))

    r = 4
    unexp_issues = [i for i in qc_result.issues if i.check_id == CheckID.CHK_UNEXPECTED_FEATURE]
    for ui in unexp_issues:
        vals = [ui.layer_name, ui.expected_value, ui.actual_value, int(ui.measurement or 0), ui.details]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_unexp.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 3:
                cell.fill = fill_error
        r += 1

    autofit_columns(ws_unexp)

    # ==========================================
    # 8. Sheet: Closure Issues
    # ==========================================
    ws_clos = wb.create_sheet(title="Closure Issues")
    ws_clos.views.sheetView[0].showGridLines = True

    ws_clos["A1"] = "POLYLINE CLOSURE ISSUES"
    ws_clos["A1"].font = title_font

    clos_headers = ["Layer Name", "Feature OID", "Gap Distance (m)", "Tolerance (m)", "Location (X, Y)", "Details"]
    for c_idx, h in enumerate(clos_headers, 1):
        ws_clos.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_clos, 3, len(clos_headers))

    r = 4
    clos_issues = [i for i in qc_result.issues if i.check_id == CheckID.CHK_CLOSURE]
    for ci in clos_issues:
        loc_str = f"({ci.location[0]:.3f}, {ci.location[1]:.3f})" if ci.location else "N/A"
        vals = [ci.layer_name, ci.object_id or "Layer Summary", ci.measurement or 0.0, ci.threshold or 0.0, loc_str, ci.details]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_clos.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 3 and isinstance(val, (int, float)) and val > 0:
                cell.fill = fill_error
        r += 1

    autofit_columns(ws_clos)

    # ==========================================
    # 9. Sheet: Feature Count Differences
    # ==========================================
    ws_cnt = wb.create_sheet(title="Feature Count Differences")
    ws_cnt.views.sheetView[0].showGridLines = True

    ws_cnt["A1"] = "FEATURE COUNT DIFFERENCES"
    ws_cnt["A1"].font = title_font

    cnt_headers = ["Layer Name", "Received Count", "Delta", "% Difference", "Details"]
    for c_idx, h in enumerate(cnt_headers, 1):
        ws_cnt.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_cnt, 3, len(cnt_headers))

    r = 4
    cnt_issues = [i for i in qc_result.issues if i.check_id == CheckID.CHK_FEATURE_COUNT]
    for ci in cnt_issues:
        ref_c = int(ci.expected_value or 0)
        rec_c = int(ci.actual_value or 0)
        delta = rec_c - ref_c
        pct = (delta / ref_c * 100.0) if ref_c > 0 else 0.0
        vals = [ci.layer_name, rec_c, delta, f"{pct:+.1f}%", ci.details]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_cnt.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 3:
                cell.fill = fill_warning
        r += 1

    autofit_columns(ws_cnt)

    # ==========================================
    # 10. Sheet: Geometry QC Summary
    # ==========================================
    ws_gsum = wb.create_sheet(title="Geometry QC Summary")
    ws_gsum.views.sheetView[0].showGridLines = True

    ws_gsum["A1"] = "GEOMETRY QC SUMMARY (10 Checks)"
    ws_gsum["A1"].font = title_font

    gsum_headers = ["Check ID", "Check Name", "Issues Found", "Status"]
    for c_idx, h in enumerate(gsum_headers, 1):
        ws_gsum.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_gsum, 3, len(gsum_headers))

    r = 4
    if qc_result.geometry_qc_summary:
        for chk_id, s_info in qc_result.geometry_qc_summary.items():
            vals = [chk_id, s_info.get("name", chk_id), s_info.get("count", 0), s_info.get("status", "PASS")]
            for c_idx, val in enumerate(vals, 1):
                cell = ws_gsum.cell(row=r, column=c_idx, value=val)
                cell.font = normal_font
                cell.border = thin_border
                if c_idx == 4:
                    cell.fill = fill_pass if val == "PASS" else fill_error
            r += 1

    # Per-Layer Geometry QC Breakdown Table
    if qc_result.layer_geometry_qc_summaries:
        r += 2
        ws_gsum.cell(row=r, column=1, value="GEOMETRY-AWARE ISSUES BREAKDOWN BY CAD LAYER").font = title_font
        r += 1
        lsum_headers = ["CAD Layer Name", "Geometry", "QC Profile", "Applicable Checks", "Features Analyzed", "Total Issues", "Errors", "Warnings", "Layer Status"]
        for c_idx, h in enumerate(lsum_headers, 1):
            ws_gsum.cell(row=r, column=c_idx, value=h)
        style_header_row(ws_gsum, r, len(lsum_headers))
        r += 1

        for lname, linfo in sorted(qc_result.layer_geometry_qc_summaries.items()):
            errs = linfo.get("errors", 0)
            warns = linfo.get("warnings", 0)
            tot = linfo.get("total_issues", errs + warns)
            st = "ERROR" if errs > 0 else ("WARNING" if warns > 0 else "PASS")
            g_type = linfo.get("geometry_type") or linfo.get("expected_geom", "Unknown")
            qc_prof = linfo.get("qc_profile") or "Generic QC"
            app_checks = linfo.get("applicable_checks", [])
            app_str = f"{len(app_checks)} checks" if app_checks else "Default"

            lvals = [lname, g_type, qc_prof, app_str, linfo.get("feature_count", "—"), tot, errs, warns, st]
            for c_idx, val in enumerate(lvals, 1):
                cell = ws_gsum.cell(row=r, column=c_idx, value=val)
                cell.font = normal_font
                cell.border = thin_border
                if c_idx == 9:
                    cell.fill = fill_pass if st == "PASS" else (fill_error if st == "ERROR" else fill_warning)
            r += 1

    autofit_columns(ws_gsum)

    # ==========================================
    # 11. Sheet: Geometry QC Issues
    # ==========================================
    ws_giss = wb.create_sheet(title="Geometry QC Issues")
    ws_giss.views.sheetView[0].showGridLines = True

    ws_giss["A1"] = "GEOMETRY & TOPOLOGY QC ISSUES"
    ws_giss["A1"].font = title_font

    g_issues = qc_result.geometry_qc_issues()

    giss_headers = ["Issue ID", "Check ID", "Issue Type", "Severity", "CAD Layer", "Geometry Type", "Feature OID", "Related OID", "Measurement", "Unit", "Location (X, Y)", "Details"]
    for c_idx, h in enumerate(giss_headers, 1):
        ws_giss.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_giss, 3, len(giss_headers))

    r = 4
    for gi in g_issues:
        loc_str = f"({gi.location[0]:.3f}, {gi.location[1]:.3f})" if gi.location else "N/A"
        vals = [
            gi.issue_id,
            gi.check_id,
            gi.issue_type,
            gi.severity,
            gi.layer_name or "N/A",
            gi.geometry_type or "N/A",
            gi.object_id,
            gi.related_object_id,
            gi.measurement,
            gi.unit,
            loc_str,
            gi.details,
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_giss.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 4:
                cell.fill = fill_error if val == Severity.ERROR else fill_warning
        r += 1

    autofit_columns(ws_giss)

    # ==========================================
    # 12. Sheet: Issue Details
    # ==========================================
    ws_all = wb.create_sheet(title="Issue Details")
    ws_all.views.sheetView[0].showGridLines = True

    ws_all["A1"] = "COMPLETE QC ISSUE REGISTRY"
    ws_all["A1"].font = title_font

    all_headers = ["Issue ID", "Check ID", "Issue Type", "Severity", "Source", "Layer Name", "Geometry Type", "OID", "Related OID", "Property", "Expected", "Actual", "Measurement", "Unit", "Location", "Details"]
    for c_idx, h in enumerate(all_headers, 1):
        ws_all.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_all, 3, len(all_headers))

    r = 4
    for iss in qc_result.issues:
        loc_str = f"({iss.location[0]:.3f}, {iss.location[1]:.3f})" if iss.location else "N/A"
        vals = [
            iss.issue_id,
            iss.check_id,
            iss.issue_type,
            iss.severity,
            iss.source,
            iss.layer_name,
            iss.geometry_type or "N/A",
            iss.object_id,
            iss.related_object_id,
            iss.property_name,
            iss.expected_value,
            iss.actual_value,
            iss.measurement,
            iss.unit,
            loc_str,
            iss.details,
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws_all.cell(row=r, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if c_idx == 4:
                if val == Severity.ERROR:
                    cell.fill = fill_error
                elif val == Severity.WARNING:
                    cell.fill = fill_warning
                elif val == Severity.INFO:
                    cell.fill = fill_info
        r += 1

    autofit_columns(ws_all)

    # Ensure target directory exists and save
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return os.path.abspath(output_path)
